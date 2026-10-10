"""Build summary tables and the downloadable Excel report."""
from datetime import datetime
from io import BytesIO

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from build import Build, estimate_power_draw, total_ram_gb, total_storage_gb

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def parts_table(build: Build) -> pd.DataFrame:
    cpu, gpu = build.cpu, build.gpu
    rows = [
        ("Processor (CPU)", cpu["name"],
         f"{cpu['cores']} cores, up to {cpu['boost_ghz']} GHz, socket {cpu['socket']}, "
         f"{cpu['power_w']} W max"),
        ("Motherboard", build.motherboard["name"],
         f"{build.motherboard['socket']}, {build.motherboard['ram_type']}, "
         f"{build.motherboard['form_factor']}, {build.motherboard['ram_slots']} RAM slots"),
        ("Memory (RAM)", f"{build.ram_kits} x {build.ram['name']}",
         f"{total_ram_gb(build)} GB {build.ram['ram_type']} total"),
    ]
    for i, drive in enumerate(build.storage, start=1):
        rows.append((f"Storage {i}", drive["name"], f"{drive['type']}, {drive['capacity_gb']} GB"))
    rows += [
        ("Graphics card (GPU)", gpu["name"],
         f"Tier {gpu['gpu_level']}/12, {gpu['power_w']} W" if gpu["power_w"]
         else "Integrated graphics"),
        ("CPU cooler", build.cooler["name"], f"~{build.cooler['capacity_w']} W cooling capacity"),
        ("Case", build.case["name"],
         f"Up to {build.case['max_form_factor']}, {build.case['airflow'].lower()} airflow"),
        ("Extra case fans", str(build.case_fans), ""),
        ("Power supply (PSU)", build.psu["name"],
         f"{build.psu['watts']} W for ~{estimate_power_draw(build)} W peak draw"),
        ("Operating system", build.os_name, build.os_family),
        ("PC age", f"{build.age_years:g} years", ""),
    ]
    return pd.DataFrame(rows, columns=["Component", "Selected part", "Key specs"])


def _style_sheet(ws, header_row=1):
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(width + 3, 70)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def build_excel(build: Build, score, level, rmse, model_name, features,
                check, temperature_c, fps_df) -> bytes:
    """Return an .xlsx report with the parts, predicted score, model inputs and FPS."""
    summary = pd.DataFrame([
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Predicted Performance Score", f"{score} / 1000" if score is not None else "n/a"),
        ("Performance level", level or "n/a"),
        ("Typical prediction error", f"± {rmse:.0f} points"),
        ("ML model", model_name),
        ("Compatibility", "All parts compatible" if check.ok else "; ".join(check.errors)),
        ("Warnings", "; ".join(check.warnings) or "None"),
        ("Estimated peak power", f"{estimate_power_draw(build)} W"),
        ("PSU capacity", f"{build.psu['watts']} W"),
        ("Estimated CPU temperature", f"{temperature_c:.0f} °C"),
        ("Total RAM", f"{total_ram_gb(build)} GB"),
        ("Total storage", f"{total_storage_gb(build)} GB"),
    ], columns=["Item", "Value"])

    inputs = pd.DataFrame([{"Model input": k, "Value": v} for k, v in features.items()])

    fps_sheet = fps_df.copy()
    fps_sheet.columns = [c if c in ("Game", "Genre") else f"{c} FPS" for c in fps_sheet.columns]

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        parts_table(build).to_excel(writer, sheet_name="Parts", index=False)
        fps_sheet.to_excel(writer, sheet_name="Game FPS", index=False)
        inputs.to_excel(writer, sheet_name="Model Inputs", index=False)
        for ws in writer.book.worksheets:
            _style_sheet(ws)
        ws = writer.book["Game FPS"]
        ws.cell(row=ws.max_row + 2, column=1,
                value="Estimated FPS at High settings, native resolution, no upscaling. "
                      "Rule-based estimate from GPU/CPU gaming ratings, not the ML model.")
    return buffer.getvalue()
